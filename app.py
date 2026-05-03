"""
Lab-first timetable generator backend with Supabase persistence hooks.
Supports multi-year generation with shared professor conflict detection.
"""

import os
import json
import random
import time
import uuid
from typing import Dict, List, Any, Optional, Tuple, Set

from dotenv import load_dotenv
load_dotenv()

from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS

DEBUG = True

def log(msg):
    if DEBUG:
        print(msg)

try:
    from supabase import create_client
except ImportError:
    create_client = None  # type: ignore

app = Flask(__name__, static_folder='static')
CORS(app)

SUPABASE_URL = os.getenv("SUPABASE_URL", "").strip()
SUPABASE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY", "").strip()

# ─────────────────────────────────────────────────────────
# Serve Frontend
# ─────────────────────────────────────────────────────────

@app.route('/')
def index():
    return send_from_directory('static', 'index.html')

@app.route('/static/<path:path>')
def serve_static(path):
    return send_from_directory('static', path)

# ─────────────────────────────────────────────────────────
# Scheduling Engine
# ─────────────────────────────────────────────────────────

class TimetableGenerator:
    def __init__(self, data: Dict[str, Any], shared_teacher_schedule: Optional[Dict[str, Dict[Tuple[int, int], str]]] = None):
        self.departments = data.get('departments', [])
        self.num_days = int(data.get('numDays', 5))
        self.num_slots = int(data.get('numSlots', 6))
        self.schedule_mode = data.get('scheduleMode', 'class')
        self.sections_per_cluster = int(data.get('sectionsPerCluster', 2))
        self.constraints = data.get('constraints', {})
        self.subjects = data.get('subjects', [])
        self.labs = data.get('labs', [])
        self.lab_assignments = data.get('labAssignments', [])
        self.classrooms = data.get('classrooms', {})
        self.teachers = data.get('teachers', [])
        # Backward-compat (old model)
        self.subject_professor_map = data.get('subjectProfessorMap', {})
        # New model: (section, subject) -> teacher_id
        self.section_subject_teacher_map = data.get('sectionSubjectTeacherMap', {})
        self.retry_seed = data.get('retrySeed')
        self.rng = random.Random(self.retry_seed if self.retry_seed is not None else time.time_ns())
        
        # Shared teacher schedule for multi-year conflict prevention
        self.shared_teacher_schedule = shared_teacher_schedule
        
        # Finalize subject type support
        for s in self.subjects:
            if 'type' not in s:
                s['type'] = 'THEORY'
                
        # teacher_id -> teacher object
        self.teacher_map = {}
        for t in self.teachers:
            teacher_id = str(t.get('id', '')).strip()
            teacher_name = str(t.get('name', '')).strip()
            if teacher_id:
                self.teacher_map[teacher_id] = t
            elif teacher_name:
                self.teacher_map[teacher_name] = t
        self.subject_map = {str(s.get('name', '')).strip(): s for s in self.subjects}
        self.teacher_unavailable = self._build_teacher_unavailability()
        self.section_cluster_map: Dict[str, int] = {}
        self.sections = self._build_sections()

    def _get_mapped_teacher_id(self, section: str, subject: str) -> str:
        subject_s = str(subject).strip()
        key = f"{section}::{subject_s}"
        if key in self.section_subject_teacher_map:
            return str(self.section_subject_teacher_map.get(key, '')).strip()
        if self.schedule_mode != 'cluster':
            return str(self.subject_professor_map.get(subject_s, '')).strip()
        cluster = self.section_cluster_map.get(section, 1)
        legacy_key = f"{subject_s}::{cluster}"
        return str(self.subject_professor_map.get(legacy_key, '')).strip()

    def _get_teacher_display(self, teacher_id: str) -> str:
        teacher = self.teacher_map.get(teacher_id) or {}
        name = str(teacher.get('name', '')).strip()
        if teacher_id and teacher_id != name:
            return f"{name} [{teacher_id}]".strip()
        return name or teacher_id

    def _build_sections(self) -> List[str]:
        sections: List[str] = []
        global_idx: int = 0
        for dept in self.departments:
            dept_name = str(dept.get('name', '')).strip()
            count = int(dept.get('sections', 1))
            for i in range(count):
                section = f"{dept_name} - Section {chr(65 + i)}"
                sections.append(section)
                if self.schedule_mode == 'cluster':
                    cluster = (global_idx // max(self.sections_per_cluster, 1)) + 1
                else:
                    cluster = 1
                self.section_cluster_map[section] = cluster
                global_idx += 1
        return sections

    def _build_teacher_unavailability(self) -> Dict[str, Set[Tuple[int, int]]]:
        unavailable: Dict[str, Set[Tuple[int, int]]] = {}
        for teacher in self.teachers:
            teacher_id = str(teacher.get('id', '')).strip()
            name = str(teacher.get('name', '')).strip()
            key = teacher_id if teacher_id else name
            unavailable[key] = set()
            for entry in teacher.get('availability', []):
                day = int(entry.get('day', -1))
                slot = int(entry.get('slot', -1))
                if 0 <= day < self.num_days and 0 <= slot < self.num_slots:
                    unavailable[key].add((day, slot))
        return unavailable

    def _empty_grid(self) -> List[List[Optional[Dict[str, Any]]]]:
        return [[None for _ in range(self.num_slots)] for _ in range(self.num_days)]

    def _validate_and_place_labs(
        self,
        grids: Dict[str, List[List[Optional[Dict[str, Any]]]]],
        teacher_schedule: Dict[str, Dict[Tuple[int, int], str]],
        room_schedule: Dict[str, Dict[Tuple[int, int], str]],
    ) -> List[str]:
        """Place lab assignments into grids; skip (with warning) any that conflict."""
        lab_warnings: List[str] = []

        for assignment in self.lab_assignments:
            section = str(assignment.get('section', '')).strip()
            subject = str(assignment.get('subject', '')).strip()
            room = str(assignment.get('room', '')).strip()
            day = int(assignment.get('day', -1))
            start_slot = int(assignment.get('slot', -1))
            duration = int(assignment.get('duration', 2))
            teacher_id = self._get_mapped_teacher_id(section, subject)
            teacher_display = self._get_teacher_display(teacher_id) if teacher_id else ''
            label = f"{subject} in {section} (Day {day+1}, Slot {start_slot+1})"

            skip_reason = None
            if section not in grids:
                skip_reason = f"Invalid section: {section}"
            elif subject not in self.subject_map:
                skip_reason = f"Invalid lab subject: {subject}"
            elif self.subject_map[subject].get('type', 'THEORY') != 'LAB':
                skip_reason = f"Subject {subject} is not a LAB type"
            elif day < 0 or day >= self.num_days:
                skip_reason = f"Invalid day ({day}) for {subject}"
            elif start_slot < 0 or (start_slot + duration) > self.num_slots:
                skip_reason = f"Lab duration exceeds available slots for {subject}"
            elif not teacher_id:
                skip_reason = f"No professor mapped for {subject} in {section}"
            elif self.schedule_mode == 'cluster' and not self._teacher_allowed_for_section(teacher_id, section):
                skip_reason = f"Cluster isolation blocked {teacher_display} for {section}"

            if skip_reason:
                log(f"[LAB-SKIP] {label} → {skip_reason}")
                lab_warnings.append(f"Skipped lab {label}: {skip_reason}")
                continue

            conflict = None
            for slot in range(start_slot, start_slot + duration):
                if grids[section][day][slot] is not None:
                    conflict = f"Class overlap at Day {day+1} Slot {slot+1} in {section}"
                    break
                if room_schedule.get(room, {}).get((day, slot)):
                    occupier = room_schedule[room][(day, slot)]
                    conflict = f"Room {room} already occupied at Day {day+1} Slot {slot+1} by {occupier}"
                    break
                if teacher_id and teacher_schedule.get(teacher_id, {}).get((day, slot)):
                    occupier = teacher_schedule[teacher_id][(day, slot)]
                    conflict = f"Professor {teacher_display} already teaching at Day {day+1} Slot {slot+1} ({occupier})"
                    break
                if teacher_id and (day, slot) in self.teacher_unavailable.get(teacher_id, set()):
                    conflict = f"Teacher {teacher_display} unavailable at Day {day+1} Slot {slot+1}"
                    break

            if conflict:
                log(f"[LAB-SKIP] {label} → {conflict}")
                lab_warnings.append(f"Skipped lab {label}: {conflict}")
                continue

            log(f"[LAB-PLACE] {label}")
            cell = {
                'subject': subject,
                'teacher': teacher_display,
                'room': room,
                'isLab': True,
                'locked': True,
            }
            for slot in range(start_slot, start_slot + duration):
                grids[section][day][slot] = cell
                room_schedule.setdefault(room, {})[(day, slot)] = section
                if teacher_id:
                    teacher_schedule.setdefault(teacher_id, {})[(day, slot)] = section

        return lab_warnings

    def _teacher_allowed_for_section(self, teacher_name: str, section: str) -> bool:
        if self.schedule_mode != 'cluster':
            return True
        teacher_cluster = int(self.teacher_map.get(teacher_name, {}).get('cluster', 1))
        return teacher_cluster == self.section_cluster_map.get(section, 1)

    def _teacher_max_consecutive(self, teacher_name: str) -> int:
        default_limit = int(self.constraints.get('maxConsecutiveClasses', 2))
        teacher_limit = self.teacher_map.get(teacher_name, {}).get('max_consecutive_classes')
        if teacher_limit is None:
            return default_limit
        return int(teacher_limit)

    def _teacher_consecutive_count(
        self,
        teacher_schedule: Dict[str, Dict[Tuple[int, int], str]],
        teacher_name: str,
        day: int,
        slot: int
    ) -> int:
        occupied = teacher_schedule.get(teacher_name, {})
        count = 1
        s = slot - 1
        while (day, s) in occupied:
            count += 1
            s -= 1
        s = slot + 1
        while (day, s) in occupied:
            count += 1
            s += 1
        return count

    def _fit_theory_subjects(
        self,
        grids: Dict[str, List[List[Optional[Dict[str, Any]]]]],
        teacher_schedule: Dict[str, Dict[Tuple[int, int], str]],
    ) -> Tuple[bool, List[str]]:
        warnings: List[str] = []
        theory_units: List[Dict[str, str]] = []
        remaining_by_section: Dict[str, int] = {section: 0 for section in self.sections}

        for section in self.sections:
            room = str(self.classrooms.get(section, '')).strip()
            if not room:
                return False, [f"No classroom mapped for section {section}"]

            for subject in self.subjects:
                if subject.get('type', 'THEORY') != 'THEORY':
                    continue
                subject_name = str(subject.get('name', '')).strip()
                teacher_id = self._get_mapped_teacher_id(section, subject_name)
                teacher_display = self._get_teacher_display(teacher_id) if teacher_id else ''
                if not teacher_id:
                    if self.schedule_mode == 'cluster':
                        return False, [
                            f"No professor mapped for {subject_name} for section {section}"
                        ]
                    return False, [f"No professor mapped for {subject_name} for section {section}"]
                if not self._teacher_allowed_for_section(teacher_id, section):
                    return False, [f"Cluster isolation blocked {teacher_display} for {section}"]
                hours = int(subject.get('hours', 0))
                for _ in range(hours):
                    theory_units.append({
                        'section': section,
                        'subject': subject_name,
                        'teacher': teacher_id,
                        'teacherDisplay': teacher_display,
                        'room': room,
                    })
                    remaining_by_section[section] += 1

        def can_fill_first_periods() -> bool:
            if not self.constraints.get('enforceFirstPeriod', True):
                return True
            for section in self.sections:
                empty_first = 0
                for day in range(self.num_days):
                    if grids[section][day][0] is None:
                        empty_first += 1
                remaining = remaining_by_section.get(section, 0)
                # If we don't have enough remaining units to fill the empty first periods, fail
                if remaining < empty_first:
                    return False
            return True

        def get_candidates(unit: Dict[str, str]) -> List[Tuple[int, int]]:
            section = unit['section']
            teacher = unit['teacher']
            subject = unit['subject']
            candidates: List[Tuple[int, int]] = []
            for day in range(self.num_days):
                for slot in range(self.num_slots):
                    if grids[section][day][slot] is not None:
                        continue
                    if (day, slot) in self.teacher_schedule_lookup(teacher_schedule, teacher):
                        continue
                    if (day, slot) in self.teacher_unavailable.get(teacher, set()):
                        continue
                    consecutive = self._teacher_consecutive_count(teacher_schedule, teacher, day, slot)
                    if consecutive > self._teacher_max_consecutive(teacher):
                        continue
                    candidates.append((day, slot))

            # Sort: prioritize first-period slots, then days with fewer classes for this subject (spread),
            # then by slot index to fill earlier slots first
            def candidate_score(ds):
                d, s = ds
                # Count how many times this subject already placed on this day for this section
                same_subject_on_day = sum(
                    1 for sl in range(self.num_slots)
                    if grids[section][d][sl] is not None and grids[section][d][sl].get('subject') == subject
                )
                # Count total slots filled per day (spread across days)
                filled_on_day = sum(1 for sl in range(self.num_slots) if grids[section][d][sl] is not None)
                return (0 if s == 0 else 1, same_subject_on_day, filled_on_day, d, s)

            candidates.sort(key=candidate_score)
            return candidates

        _bt_start = time.time()
        _bt_attempts = [0]
        _bt_timed_out = [False]
        _BT_TIME_LIMIT = 10.0
        _BT_ATTEMPT_LIMIT = 200000

        def backtrack(remaining_units: List[Dict[str, str]]) -> bool:
            if _bt_timed_out[0]:
                return False
            _bt_attempts[0] += 1
            if _bt_attempts[0] > _BT_ATTEMPT_LIMIT or (time.time() - _bt_start) > _BT_TIME_LIMIT:
                log("[TIMEOUT] Backtracking stopped due to time/attempt limit")
                _bt_timed_out[0] = True
                return False
            if not remaining_units:
                return can_fill_first_periods()
            if not can_fill_first_periods():
                return False

            chosen_idx = -1
            chosen_candidates: List[Tuple[int, int]] = []
            for idx, unit in enumerate(remaining_units):
                cand = get_candidates(unit)
                if not cand:
                    return False
                if chosen_idx == -1 or len(cand) < len(chosen_candidates):
                    chosen_idx = idx
                    chosen_candidates = cand
                    if len(chosen_candidates) == 1:
                        break

            chosen_unit = remaining_units[chosen_idx]
            section = chosen_unit['section']
            subject = chosen_unit['subject']
            teacher_id = chosen_unit['teacher']
            teacher_display = chosen_unit.get('teacherDisplay', teacher_id)
            room = chosen_unit['room']

            next_remaining = remaining_units[:chosen_idx] + remaining_units[chosen_idx + 1:]
            self.rng.shuffle(chosen_candidates)
            for day, slot in chosen_candidates:
                if _bt_timed_out[0]:
                    break
                grids[section][day][slot] = {
                    'subject': subject,
                    'teacher': teacher_display,
                    'room': room,
                    'isLab': False,
                }
                log(f"[PLACE] {subject} → Day {day} Slot {slot}")
                teacher_schedule.setdefault(teacher_id, {})[(day, slot)] = section
                remaining_by_section[section] -= 1

                if backtrack(next_remaining):
                    return True

                log(f"[BACKTRACK] Removing {subject} from Day {day} Slot {slot}")
                grids[section][day][slot] = None
                teacher_schedule.get(teacher_id, {}).pop((day, slot), None)
                if not teacher_schedule.get(teacher_id):
                    teacher_schedule.pop(teacher_id, None)
                remaining_by_section[section] += 1
            return False

        self.rng.shuffle(theory_units)
        log("[INFO] Starting backtracking")
        success = backtrack(theory_units)
        log(f"[INFO] Backtracking finished mapping, success={success}")

        if not success:
            log("[INFO] Switching to greedy fallback")
            warnings.append("Strict scheduling failed; using relaxed greedy placement")

            # Clear all theory placements from the failed backtracking attempt
            for section in self.sections:
                for day in range(self.num_days):
                    for slot in range(self.num_slots):
                        cell = grids[section][day][slot]
                        if cell is not None and not cell.get('isLab') and not cell.get('locked'):
                            grids[section][day][slot] = None
            # Remove any placements for current year's sections from the shared teacher schedule
            for tid in list(teacher_schedule.keys()):
                for ds in list(teacher_schedule[tid].keys()):
                    if teacher_schedule[tid][ds] in self.sections:
                        teacher_schedule[tid].pop(ds)
                if not teacher_schedule[tid]:
                    teacher_schedule.pop(tid)
                    
            # Re-add lab and locked placements for the current year
            for section in self.sections:
                for day in range(self.num_days):
                    for slot in range(self.num_slots):
                        cell = grids[section][day][slot]
                        if cell and (cell.get('isLab') or cell.get('locked')):
                            teacher_display = cell.get('teacher', '')
                            for tid, tobj in self.teacher_map.items():
                                t_disp = self._get_teacher_display(tid)
                                if t_disp == teacher_display:
                                    teacher_schedule.setdefault(tid, {})[(day, slot)] = section
                                    break

            # Recalculate remaining
            for section in self.sections:
                remaining_by_section[section] = sum(
                    1 for u in theory_units if u['section'] == section
                )

            # Sort units: place subjects with fewer candidates first (most constrained first)
            self.rng.shuffle(theory_units)
            theory_units.sort(key=lambda u: len(get_candidates(u)))

            unplaced: List[Dict[str, str]] = []
            for unit in theory_units:
                subject = unit['subject']
                candidates = get_candidates(unit)
                if candidates:
                    day, slot = candidates[0]
                    section = unit['section']
                    teacher_id = unit['teacher']
                    teacher_display = unit.get('teacherDisplay', teacher_id)
                    room = unit['room']
                    log(f"[FALLBACK-PLACE] {subject} → {section} Day {day} Slot {slot}")
                    grids[section][day][slot] = {
                        'subject': subject,
                        'teacher': teacher_display,
                        'room': room,
                        'isLab': False,
                    }
                    teacher_schedule.setdefault(teacher_id, {})[(day, slot)] = section
                    remaining_by_section[section] -= 1
                else:
                    log(f"[FALLBACK-SKIP] {subject} → No valid slot found")
                    unplaced.append(unit)

            if unplaced:
                grouped: Dict[str, int] = {}
                for u in unplaced:
                    key = f"{u['subject']} ({u['section']})"
                    grouped[key] = grouped.get(key, 0) + 1
                for key, count in grouped.items():
                    warnings.append(f"Could not place {count} hour(s) of {key}")

            placed_count = len(theory_units) - len(unplaced)
            if placed_count == 0 and len(theory_units) > 0:
                return False, warnings

        return True, warnings

    def teacher_schedule_lookup(self, teacher_schedule: Dict[str, Dict[Tuple[int, int], str]], teacher: str) -> Set[Tuple[int, int]]:
        return set(teacher_schedule.get(teacher, {}).keys())

    def _validate_first_period(self, grids: Dict[str, List[List[Optional[Dict[str, Any]]]]]) -> List[str]:
        warnings: List[str] = []
        if not self.constraints.get('enforceFirstPeriod', True):
            return warnings
        for section in self.sections:
            for day in range(self.num_days):
                if grids[section][day][0] is None:
                    warnings.append(f"First period empty for {section} on day {day + 1}")
        return warnings

    def _format_section_grids(self, grids: Dict[str, List[List[Optional[Dict[str, Any]]]]]) -> Dict[str, List[List[Dict[str, Any]]]]:
        output: Dict[str, List[List[Dict[str, Any]]]] = {}
        for section, grid in grids.items():
            out_grid: List[List[Dict[str, Any]]] = []
            for day in range(self.num_days):
                row: List[Dict[str, Any]] = []
                for slot in range(self.num_slots):
                    cell = grid[day][slot]
                    if cell is None:
                        row.append({'subject': 'Free', 'teacher': '', 'room': '', 'isFree': True})
                    else:
                        row.append(cell)
                out_grid.append(row)
            output[section] = out_grid
        return output

    def _build_teacher_timetables(self, section_timetables: Dict[str, List[List[Dict[str, Any]]]]) -> Dict[str, List[List[Dict[str, Any]]]]:
        """Build a timetable view per teacher."""
        teacher_tt: Dict[str, List[List[Optional[Dict[str, Any]]]]] = {}
        for section_id, t_grid in section_timetables.items():
            for day_idx, day_data in enumerate(t_grid):
                for slot_idx, cell in enumerate(day_data):
                    if cell and cell.get('teacher'):
                        t_name = str(cell['teacher'])
                        if t_name not in teacher_tt:
                            empty_t_grid: List[List[Optional[Dict[str, Any]]]] = [[None for _ in range(self.num_slots)] for _ in range(self.num_days)]
                            teacher_tt[t_name] = empty_t_grid
                        if teacher_tt[t_name][day_idx][slot_idx] is None:
                            teacher_tt[t_name][day_idx][slot_idx] = {
                                'subject': cell['subject'],
                                'section': section_id,
                                'room': cell.get('room', ''),
                            }

        formatted: Dict[str, List[List[Dict[str, Any]]]] = {}
        for t_name, fmt_grid_in in teacher_tt.items():
            fmt_grid: List[List[Dict[str, Any]]] = []
            for day_idx in range(self.num_days):
                day_data_fmt: List[Dict[str, Any]] = []
                for slot_idx in range(self.num_slots):
                    cell = fmt_grid_in[day_idx][slot_idx]
                    if cell is None:
                        day_data_fmt.append({'subject': 'Free', 'section': '', 'room': '', 'isFree': True})
                    else:
                        day_data_fmt.append(cell)
                fmt_grid.append(day_data_fmt)
            formatted[t_name] = fmt_grid

        return formatted

    def generate(self) -> Dict[str, Any]:
        section_grids = {section: self._empty_grid() for section in self.sections}
        # Use shared teacher schedule if provided (for multi-year), otherwise fresh
        teacher_schedule: Dict[str, Dict[Tuple[int, int], str]] = self.shared_teacher_schedule if self.shared_teacher_schedule is not None else {}
        room_schedule: Dict[str, Dict[Tuple[int, int], str]] = {}

        lab_warnings = self._validate_and_place_labs(section_grids, teacher_schedule, room_schedule)
        success, warnings = self._fit_theory_subjects(section_grids, teacher_schedule)
        warnings = lab_warnings + warnings
        if not success:
            raise ValueError("; ".join(warnings))
        warnings.extend(self._validate_first_period(section_grids))

        section_timetables = self._format_section_grids(section_grids)
        teacher_timetables = self._build_teacher_timetables(section_timetables)
        return {
            'sectionTimetables': section_timetables,
            'teacherTimetables': teacher_timetables,
            'warnings': warnings,
            '_teacher_schedule': teacher_schedule,  # pass back for chaining
        }


class SupabaseRepository:
    def __init__(self) -> None:
        self.enabled = bool(SUPABASE_URL and SUPABASE_KEY and create_client)
        self.client = create_client(SUPABASE_URL, SUPABASE_KEY) if self.enabled else None

    def _stable_uuid(self, prefix: str, *parts: str) -> str:
        base = f"{prefix}|{'|'.join(parts)}"
        return str(uuid.uuid5(uuid.NAMESPACE_URL, base))

    def _upsert(self, table: str, rows: List[Dict[str, Any]]) -> None:
        if not self.enabled or not self.client or not rows:
            return
        try:
            self.client.table(table).upsert(rows).execute()
        except Exception:
            return

    def save_subjects(self, subjects: List[Dict[str, Any]]) -> None:
        rows = []
        for s in subjects:
            name = str(s.get('name', '')).strip()
            if not name:
                continue
            rows.append({
                'id': self._stable_uuid('subjects', name),
                'name': name,
                'type': str(s.get('type', 'THEORY')).lower(),
                'duration': int(s.get('duration', s.get('lab_duration', 1))),
            })
        self._upsert('subjects', rows)

    def save_teachers(self, teachers: List[Dict[str, Any]], default_max_consecutive: int = 2) -> None:
        rows = []
        for t in teachers:
            name = str(t.get('name', '')).strip()
            if not name:
                continue
            rows.append({
                'id': self._stable_uuid('teachers', name),
                'name': name,
                'subjects': t.get('subjects', []),
                'cluster': str(t.get('cluster', '1')),
                'max_consecutive_classes': int(t.get('max_consecutive_classes', default_max_consecutive)),
                'availability': t.get('availability', []),
            })
        self._upsert('teachers', rows)

    def save_labs(self, labs: List[Dict[str, Any]]) -> None:
        rows = []
        for lb in labs:
            name = str(lb.get('lab_name', lb.get('name', ''))).strip()
            room_number = str(lb.get('room_number', '')).strip()
            if not name or not room_number:
                continue
            rows.append({
                'id': self._stable_uuid('labs', name, room_number),
                'name': name,
                'room_number': room_number,
            })
        self._upsert('labs', rows)

    def save_classrooms(self, classrooms: Dict[str, str]) -> None:
        rows = []
        for section, room_number in classrooms.items():
            section_s = str(section).strip()
            room_s = str(room_number).strip()
            if not section_s or not room_s:
                continue
            rows.append({
                'section_id': section_s,
                'room_number': room_s,
            })
        self._upsert('classrooms', rows)

    def save_timetable(self, entries: List[Dict[str, Any]]) -> None:
        rows = []
        for e in entries:
            class_id = str(e.get('class_id', '')).strip()
            day = str(e.get('day', '')).strip()
            slot = int(e.get('slot', 0))
            if not class_id or not day:
                continue
            rows.append({
                'class_id': class_id,
                'day': int(e.get('day', 0)),
                'slot': slot,
                'subject': str(e.get('subject', '')).strip(),
                'teacher': str(e.get('teacher', '')).strip(),
                'room': str(e.get('room', '')).strip(),
            })
        self._upsert('timetable', rows)

    def persist_input_payload(self, data: Dict[str, Any]) -> None:
        """Persist the entire input configuration to Supabase."""
        self.save_subjects(data.get('subjects', []))
        constraints = data.get('constraints', {})
        default_max_consecutive = int(constraints.get('maxConsecutiveClasses', 2))
        self.save_teachers(data.get('teachers', []), default_max_consecutive)
        self.save_labs(data.get('labs', []))
        self.save_classrooms(data.get('classrooms', {}))

    def persist_generated_timetable(self, result: Dict[str, Any]) -> None:
        """Flatten and persist the generated timetable to Supabase."""
        section_timetables = result.get('sectionTimetables', {})
        entries = []
        for section_id, grid in section_timetables.items():
            for day_idx, day_data in enumerate(grid):
                for slot_idx, cell in enumerate(day_data):
                    if not cell or cell.get('isFree') or cell.get('isLunch'):
                        continue
                    entries.append({
                        'class_id': section_id,
                        'day': day_idx,
                        'slot': slot_idx,
                        'subject': cell.get('subject', ''),
                        'teacher': cell.get('teacher', ''),
                        'room': cell.get('room', ''),
                    })
        self.save_timetable(entries)

    def _select_all(self, table: str) -> List[Dict[str, Any]]:
        if not self.enabled or not self.client:
            return []
        try:
            response = self.client.table(table).select("*").execute()
            return response.data or []
        except Exception:
            return []

    def get_teachers(self) -> List[Dict[str, Any]]:
        return self._select_all('teachers')

    def get_subjects(self) -> List[Dict[str, Any]]:
        return self._select_all('subjects')

    def get_labs(self) -> List[Dict[str, Any]]:
        return self._select_all('labs')

    def get_classrooms(self) -> List[Dict[str, Any]]:
        return self._select_all('classrooms')

    def get_timetable(self) -> List[Dict[str, Any]]:
        return self._select_all('timetable')


supabase_repo = SupabaseRepository()


def save_teachers(teachers: List[Dict[str, Any]], default_max_consecutive: int = 2) -> None:
    supabase_repo.save_teachers(teachers, default_max_consecutive)


def save_subjects(subjects: List[Dict[str, Any]]) -> None:
    supabase_repo.save_subjects(subjects)


def save_labs(labs: List[Dict[str, Any]]) -> None:
    supabase_repo.save_labs(labs)


def save_classrooms(classrooms: Dict[str, str]) -> None:
    supabase_repo.save_classrooms(classrooms)


def save_timetable(entries: List[Dict[str, Any]]) -> None:
    supabase_repo.save_timetable(entries)


def get_teachers() -> List[Dict[str, Any]]:
    return supabase_repo.get_teachers()


def get_subjects() -> List[Dict[str, Any]]:
    return supabase_repo.get_subjects()


def get_labs() -> List[Dict[str, Any]]:
    return supabase_repo.get_labs()


def get_classrooms() -> List[Dict[str, Any]]:
    return supabase_repo.get_classrooms()


def get_timetable() -> List[Dict[str, Any]]:
    return supabase_repo.get_timetable()


# ─────────────────────────────────────────────────────────
# Excel Parsing Helpers
# ─────────────────────────────────────────────────────────

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None  # type: ignore


def _normalize_header(h: str) -> str:
    """Lowercase, strip, collapse whitespace to single space."""
    import re
    return re.sub(r'\s+', ' ', str(h).strip().lower())


def _find_header_row(ws, max_scan: int = 15):
    """Scan first `max_scan` rows for the likely header row.
    Returns (row_index_1_based, {normalized_header: col_index_0_based})."""
    for row_idx in range(1, max_scan + 1):
        cells = [str(c.value or '').strip() for c in ws[row_idx]]
        non_empty = [c for c in cells if c]
        if len(non_empty) >= 3:
            mapping = {}
            for col_idx, cell_val in enumerate(cells):
                if cell_val:
                    mapping[_normalize_header(cell_val)] = col_idx
            return row_idx, mapping
    return None, {}


# Column alias maps — maps our canonical key to possible header variants
_FACULTY_ALLOTMENT_ALIASES = {
    'faculty_id': ['faculty id', 'fac id', 'faculty code', 'emp id', 'employee id', 'id'],
    'faculty_name': ['faculty name', 'name of faculty', 'teacher name', 'professor name', 'name', 'faculty'],
    'subject_code': ['subject code', 'sub code', 'course code', 'code'],
    'subject_name': ['subject name', 'sub name', 'course name', 'subject', 'course'],
    'cluster': ['cluster', 'cluster no', 'cluster number', 'grp', 'group'],
    'year': ['year', 'yr', 'academic year', 'batch year'],
    'sections': ['assigned sections', 'sections', 'section', 'sec', 'allotted sections', 'assigned section'],
}

_SECTION_WISE_ALIASES = {
    'section': ['section name', 'section', 'sec', 'class'],
    'year': ['year', 'yr', 'academic year', 'batch year'],
    'subject_code': ['subject code', 'sub code', 'course code', 'code'],
    'subject_name': ['subject name', 'sub name', 'course name', 'subject', 'course'],
    'faculty_name': ['faculty name', 'name of faculty', 'teacher name', 'professor name', 'name', 'faculty'],
    'faculty_id': ['faculty id', 'fac id', 'faculty code', 'emp id', 'employee id', 'id'],
}


def _resolve_columns(header_map: dict, aliases: dict) -> dict:
    """Given a {normalized_header: col_idx} map and an alias dict, resolve
    canonical keys to column indices. Returns {canonical_key: col_idx}."""
    resolved = {}
    for canonical, candidates in aliases.items():
        for candidate in candidates:
            if candidate in header_map:
                resolved[canonical] = header_map[candidate]
                break
    return resolved


def _parse_faculty_allotment(ws) -> dict:
    """Parse a 'Faculty allotment Cluster-X.xlsx' style worksheet."""
    header_row, header_map = _find_header_row(ws)
    if not header_row:
        return {'error': 'Could not find header row in Faculty allotment sheet.'}

    cols = _resolve_columns(header_map, _FACULTY_ALLOTMENT_ALIASES)
    missing = [k for k in ['faculty_id', 'faculty_name', 'subject_name'] if k not in cols]
    if missing:
        return {'error': f'Missing required columns in Faculty allotment: {", ".join(missing)}. Found headers: {list(header_map.keys())}'}

    records = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=False):
        vals = [c.value for c in row]
        fac_id = str(vals[cols['faculty_id']] or '').strip()
        fac_name = str(vals[cols['faculty_name']] or '').strip()
        if not fac_id and not fac_name:
            continue  # skip empty rows

        subj_code = str(vals[cols.get('subject_code', -1)] or '').strip() if 'subject_code' in cols else ''
        subj_name = str(vals[cols['subject_name']] or '').strip()
        cluster = str(vals[cols.get('cluster', -1)] or '1').strip() if 'cluster' in cols else '1'
        year = str(vals[cols.get('year', -1)] or '').strip() if 'year' in cols else ''
        sections_raw = str(vals[cols.get('sections', -1)] or '').strip() if 'sections' in cols else ''

        # Parse cluster number
        try:
            cluster_num = int(''.join(c for c in cluster if c.isdigit()) or '1')
        except ValueError:
            cluster_num = 1

        # Parse year
        year_num = None
        if year:
            digits = ''.join(c for c in year if c.isdigit())
            if digits:
                year_num = int(digits)
                if year_num > 4:
                    year_num = None  # invalid

        # Parse sections
        section_list = []
        if sections_raw:
            import re
            section_list = [s.strip() for s in re.split(r'[,;/\n]', sections_raw) if s.strip()]

        records.append({
            'faculty_id': fac_id,
            'faculty_name': fac_name,
            'subject_code': subj_code,
            'subject_name': subj_name,
            'cluster': cluster_num,
            'year': year_num,
            'sections': section_list,
        })

    return {'type': 'faculty_allotment', 'records': records}


def _parse_section_wise(ws) -> dict:
    """Parse a 'Section wise faculty allotment.xlsx' style worksheet."""
    header_row, header_map = _find_header_row(ws)
    if not header_row:
        return {'error': 'Could not find header row in Section-wise allotment sheet.'}

    cols = _resolve_columns(header_map, _SECTION_WISE_ALIASES)
    missing = [k for k in ['section', 'subject_name', 'faculty_name'] if k not in cols]
    if missing:
        return {'error': f'Missing required columns in Section-wise allotment: {", ".join(missing)}. Found headers: {list(header_map.keys())}'}

    records = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=False):
        vals = [c.value for c in row]
        section = str(vals[cols['section']] or '').strip()
        fac_name = str(vals[cols['faculty_name']] or '').strip()
        if not section and not fac_name:
            continue

        subj_code = str(vals[cols.get('subject_code', -1)] or '').strip() if 'subject_code' in cols else ''
        subj_name = str(vals[cols['subject_name']] or '').strip()
        fac_id = str(vals[cols.get('faculty_id', -1)] or '').strip() if 'faculty_id' in cols else ''
        year_raw = str(vals[cols.get('year', -1)] or '').strip() if 'year' in cols else ''

        year_num = None
        if year_raw:
            digits = ''.join(c for c in year_raw if c.isdigit())
            if digits:
                year_num = int(digits)
                if year_num > 4:
                    year_num = None

        records.append({
            'section': section,
            'year': year_num,
            'subject_code': subj_code,
            'subject_name': subj_name,
            'faculty_name': fac_name,
            'faculty_id': fac_id,
        })

    return {'type': 'section_wise', 'records': records}


def _detect_and_parse_workbook(file_storage) -> dict:
    """Accept a file upload, detect its type, and return parsed data."""
    if load_workbook is None:
        return {'error': 'openpyxl is not installed on the server.'}

    import io
    try:
        wb = load_workbook(io.BytesIO(file_storage.read()), data_only=True)
    except Exception as e:
        return {'error': f'Could not open Excel file: {str(e)}'}

    all_results = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row, header_map = _find_header_row(ws)
        if not header_row:
            continue

        headers_lower = set(header_map.keys())

        # Heuristic: if 'section name' or 'section' is a header and 'cluster' is not → section_wise
        has_section = any(k in headers_lower for k in ['section name', 'section', 'sec', 'class'])
        has_cluster = any(k in headers_lower for k in ['cluster', 'cluster no', 'cluster number'])

        if has_section and not has_cluster:
            result = _parse_section_wise(ws)
        else:
            result = _parse_faculty_allotment(ws)

        if 'error' in result:
            result['sheet'] = sheet_name
        else:
            result['sheet'] = sheet_name
        all_results.append(result)

    if not all_results:
        return {'error': 'No valid data sheets found in the uploaded file.'}

    return {'sheets': all_results}


# ─────────────────────────────────────────────────────────
# API Endpoints
# ─────────────────────────────────────────────────────────

@app.route('/api/upload-excel', methods=['POST'])
def upload_excel():
    """Upload and parse an Excel file for faculty/subject data."""
    try:
        if 'file' not in request.files:
            return jsonify({'error': 'No file uploaded'}), 400

        file = request.files['file']
        if not file.filename:
            return jsonify({'error': 'Empty filename'}), 400

        ext = os.path.splitext(file.filename)[1].lower()
        if ext not in ('.xlsx', '.xls'):
            return jsonify({'error': f'Unsupported file type: {ext}. Please upload .xlsx files.'}), 400

        result = _detect_and_parse_workbook(file)

        if 'error' in result:
            return jsonify({'error': result['error']}), 400

        return jsonify({'success': True, 'data': result})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500


@app.route('/api/generate', methods=['POST'])
def generate_timetable():
    """Generate timetable from provided configuration."""
    try:
        data = request.get_json()
        if not data:
            return jsonify({'error': 'No data provided'}), 400

        generator = TimetableGenerator(data)
        result = generator.generate()
        # Remove internal key before sending to client
        result.pop('_teacher_schedule', None)
        supabase_repo.persist_input_payload(data)
        supabase_repo.persist_generated_timetable(result)

        return jsonify({
            'success': True,
            'data': result,
            'config': {
                'numDays': data.get('numDays', 5),
                'numSlots': data.get('numSlots', 6),
            }
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/generate-multi-year', methods=['POST'])
def generate_multi_year_timetable():
    """Generate timetables for multiple years with shared professor conflict detection."""
    try:
        data = request.get_json()
        if not data:
            return jsonify({'error': 'No data provided'}), 400

        years_data = data.get('years', {})
        if not years_data:
            return jsonify({'error': 'No year data provided'}), 400

        # Shared teacher schedule across all years to prevent professor conflicts
        shared_teacher_schedule: Dict[str, Dict[Tuple[int, int], str]] = {}
        # Shared room schedule across all years to prevent room conflicts for labs
        shared_room_schedule: Dict[str, Dict[Tuple[int, int], str]] = {}
        
        results_by_year = {}
        all_warnings = []
        
        # Sort years so we process them in order (1st, 2nd, 3rd, 4th)
        year_keys = sorted(years_data.keys(), key=lambda x: int(x) if x.isdigit() else 0)
        
        # ── Phase 1: Place ALL labs from ALL years first ──────────────
        # This ensures user-placed labs are never skipped because a
        # theory class from another year grabbed the professor's slot.
        generators: Dict[str, TimetableGenerator] = {}
        year_grids: Dict[str, Dict[str, List[List[Optional[Dict[str, Any]]]]]] = {}
        year_lab_warnings: Dict[str, List[str]] = {}
        
        for year_key in year_keys:
            year_data = years_data[year_key]
            year_label = year_data.get('yearLabel', f'Year {year_key}')
            log(f"[MULTI-YEAR] Phase 1 — placing labs for {year_label}")
            
            gen = TimetableGenerator(year_data, shared_teacher_schedule=shared_teacher_schedule)
            grids = {section: gen._empty_grid() for section in gen.sections}
            lab_warnings = gen._validate_and_place_labs(grids, shared_teacher_schedule, shared_room_schedule)
            
            generators[year_key] = gen
            year_grids[year_key] = grids
            year_lab_warnings[year_key] = lab_warnings
        
        # ── Phase 2: Generate theory for each year ────────────────────
        # Labs are already in the grids and in shared_teacher_schedule,
        # so theory placement will work around them correctly.
        for year_key in year_keys:
            year_data = years_data[year_key]
            year_label = year_data.get('yearLabel', f'Year {year_key}')
            log(f"[MULTI-YEAR] Phase 2 — placing theory for {year_label}")
            
            gen = generators[year_key]
            grids = year_grids[year_key]
            lab_warnings = year_lab_warnings[year_key]
            
            try:
                success, theory_warnings = gen._fit_theory_subjects(grids, shared_teacher_schedule)
                warnings = lab_warnings + theory_warnings
                
                if not success:
                    raise ValueError("; ".join(warnings))
                
                warnings.extend(gen._validate_first_period(grids))
                
                section_timetables = gen._format_section_grids(grids)
                teacher_timetables = gen._build_teacher_timetables(section_timetables)
                
                result = {
                    'sectionTimetables': section_timetables,
                    'teacherTimetables': teacher_timetables,
                    'warnings': warnings,
                }
                
                results_by_year[year_key] = {
                    'yearLabel': year_label,
                    'data': result,
                }
                
                if warnings:
                    for w in warnings:
                        all_warnings.append(f"[{year_label}] {w}")
                        
            except ValueError as e:
                all_warnings.append(f"[{year_label}] {str(e)}")
                results_by_year[year_key] = {
                    'yearLabel': year_label,
                    'error': str(e),
                }

        # Build a combined teacher timetable across all years
        combined_teacher_tt: Dict[str, List[List[Optional[Dict[str, Any]]]]] = {}
        num_days = int(data.get('numDays', 5))
        num_slots = int(data.get('numSlots', 6))
        
        for year_key, year_result in results_by_year.items():
            if 'error' in year_result:
                continue
            year_label = year_result['yearLabel']
            section_tts = year_result['data'].get('sectionTimetables', {})
            for section_id, grid in section_tts.items():
                for day_idx, day_data in enumerate(grid):
                    for slot_idx, cell in enumerate(day_data):
                        if cell and cell.get('teacher') and not cell.get('isFree'):
                            t_name = str(cell['teacher'])
                            if t_name not in combined_teacher_tt:
                                combined_teacher_tt[t_name] = [[None for _ in range(num_slots)] for _ in range(num_days)]
                            if combined_teacher_tt[t_name][day_idx][slot_idx] is None:
                                combined_teacher_tt[t_name][day_idx][slot_idx] = {
                                    'subject': cell['subject'],
                                    'section': f"[{year_label}] {section_id}",
                                    'room': cell.get('room', ''),
                                }

        # Format combined teacher timetables
        formatted_teacher_tt = {}
        for t_name, grid in combined_teacher_tt.items():
            fmt_grid = []
            for day_idx in range(num_days):
                day_data = []
                for slot_idx in range(num_slots):
                    cell = grid[day_idx][slot_idx]
                    if cell is None:
                        day_data.append({'subject': 'Free', 'section': '', 'room': '', 'isFree': True})
                    else:
                        day_data.append(cell)
                fmt_grid.append(day_data)
            formatted_teacher_tt[t_name] = fmt_grid

        return jsonify({
            'success': True,
            'yearResults': results_by_year,
            'combinedTeacherTimetables': formatted_teacher_tt,
            'warnings': all_warnings,
            'config': {
                'numDays': num_days,
                'numSlots': num_slots,
            }
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500


@app.route('/api/supabase-schema', methods=['GET'])
def supabase_schema():
    schema_sql = """
create table if not exists teachers (
  id bigserial primary key,
  name text not null unique,
  subjects jsonb default '[]'::jsonb,
  max_consecutive_classes int default 2,
  availability jsonb default '[]'::jsonb
);

create table if not exists subjects (
  id bigserial primary key,
  name text not null unique,
  type text not null check (type in ('theory','lab')),
  duration int not null default 1
);

create table if not exists labs (
  id text primary key,
  name text not null,
  room_number text not null unique
);

create table if not exists classrooms (
  section_id text primary key,
  room_number text not null
);

create table if not exists timetable (
  id bigserial primary key,
  class_id text not null,
  day int not null,
  slot int not null,
  subject text not null,
  teacher text,
  room text
);
"""
    return jsonify({'success': True, 'sql': schema_sql})


@app.route('/api/save-config', methods=['POST'])
def save_config():
    """Save configuration for future use."""
    try:
        data = request.get_json()
        config_dir = os.path.join(os.path.dirname(__file__), 'configs')
        os.makedirs(config_dir, exist_ok=True)
        filename = data.get('name', 'default') + '.json'
        filepath = os.path.join(config_dir, filename)
        with open(filepath, 'w') as f:
            json.dump(data, f, indent=2)
        return jsonify({'success': True, 'message': f'Configuration saved as {filename}'})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/load-configs', methods=['GET'])
def load_configs():
    """List saved configurations."""
    try:
        config_dir = os.path.join(os.path.dirname(__file__), 'configs')
        if not os.path.exists(config_dir):
            return jsonify({'configs': []})
        configs = []
        for f in os.listdir(config_dir):
            if f.endswith('.json'):
                filepath = os.path.join(config_dir, f)
                with open(filepath, 'r') as fh:
                    config_data = json.load(fh)
                configs.append({
                    'name': f.replace('.json', ''),
                    'data': config_data
                })
        return jsonify({'configs': configs})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/load-config/<name>', methods=['GET'])
def load_config(name):
    """Load a specific saved configuration."""
    try:
        config_dir = os.path.join(os.path.dirname(__file__), 'configs')
        filepath = os.path.join(config_dir, name + '.json')
        if not os.path.exists(filepath):
            return jsonify({'error': 'Configuration not found'}), 404
        with open(filepath, 'r') as f:
            config_data = json.load(f)
        return jsonify({'success': True, 'data': config_data})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


if __name__ == '__main__':
    os.makedirs('static', exist_ok=True)
    app.run(debug=False, port=5001)
